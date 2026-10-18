"""Builds paper/out/AgriSense_Complete_Documentation.docx: every project
document in one file. Markdown sources are converted (headings, tables,
lists, code, images); PDFs (poster, the two papers) are embedded as page
images so their layout is preserved exactly. Run after impl_paper.py,
survey_paper.py and to_pdf.ps1 so the paper PDFs exist.
"""
import os
import re
import datetime

import pymupdf
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ieee import ROOT, PAPER, _shade

OUT = os.path.join(PAPER, "out")
IMG = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_pages")
os.makedirs(IMG, exist_ok=True)
NAVY = RGBColor(0x1B, 0x3A, 0x6B)
TEXT_W = 6.5

doc = Document()
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for lvl, size in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
    hs = doc.styles[f"Heading {lvl}"]
    hs.font.name, hs.font.size, hs.font.color.rgb = "Calibri", Pt(size), NAVY
    hs.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Inches(1))


def field(paragraph, instr):
    r = paragraph.add_run()
    for kind, txt in (("begin", None), (None, instr), ("separate", None), (None, "1" if instr == "PAGE" else "Right-click → Update Field"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
        elif txt == instr:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = instr
        else:
            el = OxmlElement("w:t"); el.text = txt
        r._r.append(el)


def inline(p, text, size=None, color=None):
    text = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    text = re.sub(r"<[^>]+>", "", text)
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*|~~[^~]+~~)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = p.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x8C, 0x2F, 0x39)
        elif tok.startswith("~~"):
            r = p.add_run(tok[2:-2]); r.font.strike = True
        elif tok.startswith("*") and len(tok) > 2:
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            r = p.add_run(tok)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color


def md_table(rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(re.fullmatch(r":?-{2,}:?", x) for x in c if x)]
    n = max(len(c) for c in cells)
    t = doc.add_table(rows=len(cells), cols=n)
    t.style = "Table Grid"
    for i, row in enumerate(cells):
        for j in range(n):
            cell = t.rows[i].cells[j]
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            if i == 0:
                _shade(cell, "1B3A6B")
                inline(p, val.replace("**", ""), size=9, color=RGBColor(255, 255, 255))
                for r in p.runs:
                    r.bold = True
            else:
                if i % 2 == 0:
                    _shade(cell, "EDF1F6")
                inline(p, val, size=9)
    doc.add_paragraph()


def add_image(path, width=TEXT_W):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(path, width=Inches(width))
    return p


def md(path, shift=1):
    """Convert a Markdown file; headings are shifted down by `shift` levels."""
    base = os.path.dirname(path)
    lines = open(path, encoding="utf-8").read().splitlines()
    i, para = 0, []

    def flush():
        if para:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            inline(p, " ".join(para))
            para.clear()

    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if s.startswith("```"):
            flush()
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i]); i += 1
            t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
            c = t.rows[0].cells[0]; _shade(c, "F4F6F8")
            p = c.paragraphs[0]
            for k, cl in enumerate(code):
                r = p.add_run(cl); r.font.name = "Consolas"; r.font.size = Pt(8)
                if k < len(code) - 1:
                    r.add_break()
            doc.add_paragraph()
        elif re.match(r"^#{1,6} ", s):
            flush()
            lvl = len(s) - len(s.lstrip("#"))
            title = re.sub(r"<[^>]+>", "", s[lvl:].strip())
            doc.add_heading(title.replace("**", ""), min(lvl + shift, 4))
        elif s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i]); i += 1
            md_table(rows)
            continue
        elif re.match(r"^([-*+]|\d+[.)])\s+", s):
            flush()
            indent = (len(ln) - len(ln.lstrip())) // 2
            numbered = bool(re.match(r"^\d", s))
            p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            p.paragraph_format.left_indent = Inches(0.25 + 0.25 * indent)
            inline(p, re.sub(r"^([-*+]|\d+[.)])\s+", "", s))
        elif s.startswith(">"):
            flush()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            inline(p, s.lstrip("> "), color=RGBColor(0x44, 0x44, 0x44))
            for r in p.runs:
                r.italic = True
        elif re.fullmatch(r"-{3,}|\*{3,}", s):
            flush()
        elif re.search(r"!\[[^\]]*\]\(([^)]+)\)|<img [^>]*src=\"([^\"]+)\"", s):
            flush()
            for m in re.finditer(r"!\[[^\]]*\]\(([^)]+)\)|<img [^>]*src=\"([^\"]+)\"", s):
                src = m.group(1) or m.group(2)
                fp = os.path.normpath(os.path.join(base, src))
                if os.path.exists(fp):
                    add_image(fp, width=4.8)
        elif not s:
            flush()
        else:
            para.append(s)
        i += 1
    flush()


def pdf_pages(pdf, tag, dpi=150, width=6.0, landscape=False):
    d = pymupdf.open(pdf)
    if landscape:
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Inches(11), Inches(8.5)
        width = 9.0
    for k, page in enumerate(d):
        out = os.path.join(IMG, f"{tag}_{k + 1}.png")
        page.get_pixmap(dpi=dpi).save(out)
        # break before the image, not after it, so a full-page image never
        # pushes an empty break paragraph onto its own page
        add_image(out, width).paragraph_format.page_break_before = k > 0
    if landscape:
        s = doc.add_section(WD_SECTION.NEW_PAGE)
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Inches(8.5), Inches(11)


def part(title, intro=None, new_page=True):
    doc.add_heading(title, 1).paragraph_format.page_break_before = new_page
    if intro:
        p = doc.add_paragraph(); inline(p, intro, color=RGBColor(0x55, 0x55, 0x55))
        for r in p.runs:
            r.italic = True


# ------------------------------------------------------------------ cover
for _ in range(5):
    doc.add_paragraph()
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AgriSense AI"); r.bold = True; r.font.size = Pt(40); r.font.color.rgb = NAVY
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("AI-Powered Smart Agriculture Decision Support System"); r.font.size = Pt(16)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Complete Project Documentation"); r.font.size = Pt(14); r.italic = True
for _ in range(3):
    doc.add_paragraph()
for line, size, bold in [
        ("Tisha Jhariya · Akiko Gajbhiye · Ansh Meshram · Agya Chourasia", 12, True),
        ("Guide: [Guide Name]", 11, False),
        ("Department of Emerging Technologies (Artificial Intelligence & Data Science)", 11, False),
        ("S. B. Jain Institute of Technology, Management and Research, Nagpur", 11, False),
        (datetime.date.today().strftime("%B %Y"), 11, False)]:
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(line); r.font.size = Pt(size); r.bold = bold

doc.add_page_break()
p = doc.add_paragraph()
r = p.add_run("Contents"); r.bold = True; r.font.size = Pt(20); r.font.color.rgb = NAVY
field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')

# ------------------------------------------------------------------ parts
part("1. Project Summary", "Plain-language overview of the whole project (source: AgriSense_Project_Summary.md).")
md(os.path.join(ROOT, "AgriSense_Project_Summary.md"))

part("2. Project Poster", "Conference/exhibition poster (source: AgriSense_Poster.pdf), reproduced on the next page.")
pdf_pages(os.path.join(ROOT, "AgriSense_Poster.pdf"), "poster", dpi=200, landscape=True)

part("3. Research Paper I — Implementation Paper",
     "IEEE Access-format implementation paper (editable copy: paper/out/AgriSense_Implementation_Paper.docx).",
     new_page=False)  # the poster's landscape section already ends on a page break
pdf_pages(os.path.join(OUT, "AgriSense_Implementation_Paper.pdf"), "impl")

part("4. Research Paper II — Survey Paper",
     "IEEE Access-format survey paper (editable copy: paper/out/AgriSense_Survey_Paper.docx).")
pdf_pages(os.path.join(OUT, "AgriSense_Survey_Paper.pdf"), "survey")

part("5. Project README", "Repository overview (source: README.md).")
md(os.path.join(ROOT, "README.md"))

part("6. System Workflow", "How the whole system fits together (source: documents/WORKFLOW.md).")
md(os.path.join(ROOT, "documents", "WORKFLOW.md"))

part("7. Hardware and Firmware Guides")
md(os.path.join(ROOT, "esp32_firmware", "WIRING_GUIDE.md"))
md(os.path.join(ROOT, "esp32_firmware", "ESP32CAM_UPLOAD_GUIDE.md"))

part("8. Database Setup", "Supabase schema, row-level security and pairing (sources: documents/Supabase_Complete_Setup.md, "
     "documents/supabase_presence_and_pairing.sql).")
md(os.path.join(ROOT, "documents", "Supabase_Complete_Setup.md"))

part("9. Machine Learning")
md(os.path.join(ROOT, "ml", "README.md"))
md(os.path.join(ROOT, "documents", "DATASETS.md"))
md(os.path.join(ROOT, "ml", "artifacts", "results_and_accuracy", "README.md"))
add_image(os.path.join(ROOT, "ml", "artifacts", "results_and_accuracy", "results.png"), 6.0)
add_image(os.path.join(ROOT, "ml", "artifacts", "results_and_accuracy", "confusion_matrix_normalized.png"), 6.0)
md(os.path.join(ROOT, "serving", "DEPLOY.md"))

part("10. Plant-Disease Implementation Plan", "Phase-by-phase plan for the two-stage pipeline (source: plant-disease-implementation-plan.md).")
md(os.path.join(ROOT, "plant-disease-implementation-plan.md"))

part("11. Application Screenshots")
for title, f in [("Web — irrigation controls", "documents/images/web-controls.png"),
                 ("Web — device connection and pairing", "documents/images/web-connect.png"),
                 ("Mobile — controls, plant scan, advisor", "paper/png/app_montage.png"),
                 ("Mobile — plant scan result", "Plantscan.jpeg"),
                 ("Mobile — farm advisor", "advisor.jpeg")]:
    doc.add_heading(title, 3)
    fp = os.path.join(ROOT, *f.split("/"))
    from PIL import Image
    im = Image.open(fp)
    add_image(fp, width=min(TEXT_W, 3.0 if im.height > im.width else TEXT_W))

part("12. Research Paper Build Prompt", "The prompt used to drive the research-paper work (source: AgriSense_Paper_Prompt.md).")
md(os.path.join(ROOT, "AgriSense_Paper_Prompt.md"))

part("Appendix A. Paper Audit and Gap Report",
     "Internal audit of every claim against the repository, and the list of measurements still pending (source: paper/audit.md).")
md(os.path.join(PAPER, "audit.md"))

# footer page numbers
for s in doc.sections:
    fp = s.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if not fp.runs:
        field(fp, "PAGE")

out = os.path.join(OUT, "AgriSense_Complete_Documentation.docx")
doc.save(out)
print("wrote", out)
