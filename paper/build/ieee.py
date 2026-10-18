"""Minimal IEEE Access-style DOCX writer shared by the two papers and the
compendium. Citations are written inline as [@key] or [@k1,k2] and numbered
in order of first appearance; the reference list is formatted IEEE-style
from paper/refs.bib + paper/refs_extra.bib.
"""
import os
import re

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PAPER = os.path.dirname(HERE)
ROOT = os.path.dirname(PAPER)
EQ_DIR = os.path.join(HERE, "_eq")
os.makedirs(EQ_DIR, exist_ok=True)

FONT = "Times New Roman"
ACCENT = RGBColor(0x1B, 0x3A, 0x6B)
COL_W = 3.45   # inches, one column
PAGE_W = 7.16  # inches, text width


# ---------------------------------------------------------------- bibliography
def _fields(body):
    out, i = {}, 0
    while True:
        m = re.compile(r"(\w+)\s*=\s*\{").search(body, i)
        if not m:
            return out
        depth, j = 1, m.end()
        while depth:
            depth += {"{": 1, "}": -1}.get(body[j], 0)
            j += 1
        out[m.group(1).lower()] = body[m.end():j - 1]
        i = j


def _clean(s):
    s = re.sub(r"\\url\{([^}]*)\}", r"\1", s)
    for a, b in [(r"{\'e}", "é"), (r"{\'i}", "í"), (r"{\'u}", "ú"), (r"{\'o}", "ó"),
                 (r"{\"o}", "ö"), (r"{\"u}", "ü"), (r"{\`o}", "ò"), (r"{\v{s}}", "š"),
                 (r"{\c{c}}", "ç"), (r"{\~a}", "ã"), (r"\&", "&"), ("--", "–"),
                 (r"$\times$", "×"), ("~", " ")]:
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", re.sub(r"[{}]", "", s)).strip()


def load_bib():
    bib = {}
    for name in ("refs.bib", "refs_extra.bib"):
        txt = open(os.path.join(PAPER, name), encoding="utf-8").read()
        for m in re.finditer(r"@(\w+)\{([^,]+),(.*?)\n\}", txt, re.S):
            f = {k: _clean(v) for k, v in _fields(m.group(3)).items()}
            f["_type"] = m.group(1).lower()
            bib[m.group(2).strip()] = f
    return bib


def _authors(a):
    names = [n.strip() for n in a.split(" and ")]
    out = []
    for n in names:
        if n == "others":
            out.append("et al.")
            continue
        if "," in n:
            last, first = [x.strip() for x in n.split(",", 1)]
            ini = " ".join(p[0] + "." for p in re.split(r"[\s]+", first) if p and p[0].isalpha())
            out.append(f"{ini} {last}".strip())
        else:
            out.append(n)
    if len(out) > 6:
        out = out[:1] + ["et al."]
    if out[-1] == "et al.":
        return ", ".join(out[:-1]) + " et al."
    return out[0] if len(out) == 1 else ", ".join(out[:-1]) + ", and " + out[-1]


def fmt_ref(e):
    """Returns a list of (text, italic) runs in IEEE reference style."""
    a, t, y = _authors(e.get("author", "")), e.get("title", ""), e.get("year", "")
    typ = e["_type"]
    runs = [(f"{a}, ", False)]
    if typ == "article":
        runs += [(f"“{t},” ", False), (e.get("journal", ""), True)]
        tail = ""
        if e.get("volume"):
            tail += f", vol. {e['volume']}"
        if e.get("number"):
            tail += f", no. {e['number']}"
        if e.get("pages"):
            tail += f", {'pp.' if '–' in e['pages'] else 'Art. no.'} {e['pages']}"
        runs.append((f"{tail}, {y}", False))
    elif typ == "inproceedings":
        runs += [(f"“{t},” in ", False), (e.get("booktitle", ""), True)]
        tail = f", {y}"
        if e.get("pages"):
            tail += f", pp. {e['pages']}"
        runs.append((tail, False))
    elif typ == "book":
        runs += [(t, True)]
        tail = ""
        if e.get("series"):
            tail += f" ({e['series']})"
        if e.get("publisher"):
            tail += f". {e.get('address', '')}{': ' if e.get('address') else ''}{e['publisher']}"
        runs.append((f"{tail}, {y}", False))
    else:
        runs += [(f"“{t},” ", False), (f"{y}", False)]
        if e.get("note"):
            runs.append((f". [Online]. Available: {e['note']}", False))
    if e.get("doi"):
        runs.append((f", doi: {e['doi']}", False))
    runs.append((".", False))
    return runs


# ---------------------------------------------------------------- docx helpers
def _set_cols(section, n, space=0.25):
    sectPr = section._sectPr
    for c in sectPr.findall(qn("w:cols")):
        sectPr.remove(c)
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), str(n))
    cols.set(qn("w:space"), str(int(space * 1440)))
    sectPr.append(cols)


def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def _keep_next(p):
    p.paragraph_format.keep_with_next = True


def _page_field(run):
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


class Paper:
    def __init__(self, running_head, journal="IEEE Access"):
        self.doc = Document()
        self.bib = load_bib()
        self.order = []
        self.fig_n = self.tab_n = self.eq_n = 0
        self.h1_n = 0
        self.h2_n = 0
        st = self.doc.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(10)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(0)
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.line_spacing = 1.0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        sec.left_margin = sec.right_margin = Inches(0.67)
        sec.top_margin = Inches(0.75)
        sec.bottom_margin = Inches(0.8)
        sec.header_distance = Inches(0.35)
        ht = sec.header.add_table(1, 2, Inches(PAGE_W))
        left, right = ht.rows[0].cells
        left.width, right.width = Inches(5.9), Inches(1.26)
        r = left.paragraphs[0].add_run(running_head)
        r.font.size, r.font.name, r.italic = Pt(8), FONT, True
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = right.paragraphs[0].add_run(journal)
        r.font.size, r.font.bold, r.font.color.rgb = Pt(8), True, ACCENT
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = fp.add_run()
        r.font.size = Pt(8)
        _page_field(r)

    # ------------------------------------------------------------ text
    def _cite(self, keys):
        nums = []
        for k in [k.strip() for k in keys.split(",")]:
            if k not in self.bib:
                raise KeyError(f"missing bib key: {k}")
            if k not in self.order:
                self.order.append(k)
            nums.append(self.order.index(k) + 1)
        nums = sorted(set(nums))
        # compress runs of 3+ consecutive numbers: [3]–[6]
        parts, i = [], 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
                j += 1
            if j - i >= 2:
                parts.append(f"[{nums[i]}]–[{nums[j]}]")
            else:
                parts.extend(f"[{n}]" for n in nums[i:j + 1])
            i = j + 1
        return ", ".join(parts)

    def runs(self, p, text, size=None, color=None):
        text = re.sub(r"\[@([^\]]+)\]", lambda m: self._cite(m.group(1)), text)
        for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|\^\{[^}]+\}|_\{[^}]+\})", text):
            if not tok:
                continue
            if tok.startswith("**"):
                r = p.add_run(tok[2:-2]); r.bold = True
            elif tok.startswith("*"):
                r = p.add_run(tok[1:-1]); r.italic = True
            elif tok.startswith("^{"):
                r = p.add_run(tok[2:-1]); r.font.superscript = True
            elif tok.startswith("_{"):
                r = p.add_run(tok[2:-1]); r.font.subscript = True
            else:
                r = p.add_run(tok)
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def para(self, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=None,
             after=0):
        p = self.doc.add_paragraph()
        p.alignment = align
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(after)
        return self.runs(p, text, size)

    def bullets(self, items, numbered=False):
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.first_line_indent = Inches(-0.14)
            p.paragraph_format.space_after = Pt(1)
            p.add_run(f"{i})  " if numbered else "•  ")
            self.runs(p, it)

    # ------------------------------------------------------------ structure
    def front(self, title, authors, affils, corresponding, abstract, keywords,
              received=None, funding=None):
        d = self.doc
        if received:
            p = d.add_paragraph()
            self.runs(p, received, size=7.5)
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(title)
        r.font.size, r.bold, r.font.color.rgb = Pt(21), True, RGBColor(0x11, 0x11, 0x11)
        r.font.name = "Arial"
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        self.runs(p, authors, size=10.5)
        for a in affils:
            self.runs(d.add_paragraph(), a, size=8)
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        self.runs(p, f"Corresponding author: {corresponding}", size=8)
        if funding:
            p = d.add_paragraph(); self.runs(p, funding, size=8)
            p.paragraph_format.space_after = Pt(4)
        self._rule()
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("ABSTRACT  "); r.bold = True; r.font.size = Pt(9)
        r.font.color.rgb = ACCENT
        self.runs(p, abstract, size=9)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("INDEX TERMS  "); r.bold = True; r.font.size = Pt(9)
        r.font.color.rgb = ACCENT
        self.runs(p, keywords, size=9)
        self._rule()
        self.columns(2)

    def _rule(self):
        p = self.doc.add_paragraph()
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1B3A6B")):
            bot.set(qn(f"w:{k}"), v)
        b.append(bot)
        pPr.append(b)
        p.paragraph_format.space_after = Pt(4)

    def columns(self, n):
        sec = self.doc.add_section(WD_SECTION.CONTINUOUS)
        _set_cols(sec, n)

    def h1(self, text, numbered=True):
        self.h2_n = 0
        roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI",
                 "XII", "XIII", "XIV", "XV", "XVI"]
        if numbered:
            self.h1_n += 1
            text = f"{roman[self.h1_n - 1]}. {text.upper()}"
        else:
            text = text.upper()
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9.5), ACCENT
        _keep_next(p)

    def h2(self, text):
        self.h2_n += 1
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{chr(64 + self.h2_n)}. {text.upper()}")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8.5), ACCENT
        _keep_next(p)

    def h3(self, n, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        r = p.add_run(f"{n}) {text.upper()}")
        r.italic, r.font.size, r.font.color.rgb = True, Pt(8.5), ACCENT
        _keep_next(p)

    # ------------------------------------------------------------ equations
    def eq(self, latex, label=True):
        self.eq_n += 1
        path = os.path.join(EQ_DIR, f"{abs(hash(latex))}.png")
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${latex}$", fontsize=13)
        plt.rcParams["mathtext.fontset"] = "stix"
        fig.savefig(path, dpi=400, bbox_inches="tight", pad_inches=0.02,
                    transparent=False, facecolor="white")
        plt.close(fig)
        t = self.doc.add_table(rows=1, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        w = [Inches(COL_W - 0.45), Inches(0.45)]
        for c, wd in zip(t.rows[0].cells, w):
            c.width = wd
        p = t.rows[0].cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        from PIL import Image
        w_in = Image.open(path).width / 400 * 0.78
        p.add_run().add_picture(path, width=Inches(min(w_in, COL_W - 0.6)))
        p = t.rows[0].cells[1].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if label:
            p.add_run(f"({self.eq_n})")
        return self.eq_n

    # ------------------------------------------------------------ floats
    def figure(self, path, caption, wide=False, width=None):
        self.fig_n += 1
        if wide:
            self.columns(1)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        _keep_next(p)
        p.add_run().add_picture(path, width=Inches(width or (PAGE_W if wide else COL_W)))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run(f"FIGURE {self.fig_n}. ")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8), ACCENT
        self.runs(p, caption, size=8)
        if wide:
            self.columns(2)
        return self.fig_n

    def table(self, header, rows, caption, widths=None, wide=False, size=7.5,
              first_bold=False):
        self.tab_n += 1
        if wide:
            self.columns(1)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        _keep_next(p)
        r = p.add_run(f"TABLE {self.tab_n}. ")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(8), ACCENT
        self.runs(p, caption, size=8)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        total = PAGE_W if wide else COL_W
        widths = widths or [total / len(header)] * len(header)
        scale = total / sum(widths)
        for ri, row in enumerate([header] + rows):
            for ci, val in enumerate(row):
                c = t.rows[ri].cells[ci]
                c.width = Inches(widths[ci] * scale)
                cp = c.paragraphs[0]
                cp.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                if ri == 0:
                    _shade(c, "1B3A6B")
                    rr = cp.add_run(str(val)); rr.bold = True
                    rr.font.size = Pt(size); rr.font.color.rgb = RGBColor(255, 255, 255)
                else:
                    if ri % 2 == 0:
                        _shade(c, "EDF1F6")
                    self.runs(cp, str(val), size=size)
                    if first_bold and ci == 0:
                        for rr in cp.runs:
                            rr.bold = True
        for row in t.rows[:-1]:   # keep the table on one page where possible
            for c in row.cells:
                for cp in c.paragraphs:
                    _keep_next(cp)
        for row in t.rows:
            trPr = row._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)
        if wide:
            self.columns(2)
        return self.tab_n

    def algorithm(self, title, lines):
        self.para(f"**{title}**", indent=False, size=8.5)
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        _shade(c, "F6F8FA")
        first = True
        for i, ln in enumerate(lines, 1):
            p = c.paragraphs[0] if first else c.add_paragraph()
            first = False
            r = p.add_run(f"{i:>2}: ")
            r.font.size, r.font.name = Pt(7.5), "Consolas"
            r = p.add_run(ln)
            r.font.size, r.font.name = Pt(7.5), "Consolas"
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ------------------------------------------------------------ back matter
    def references(self):
        self.h1("References", numbered=False)
        for i, k in enumerate(self.order, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f"[{i}]\t")
            r.font.size = Pt(7.5)
            for txt, it in fmt_ref(self.bib[k]):
                r = p.add_run(txt)
                r.italic, r.font.size = it, Pt(7.5)

    def bio(self, name, text):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(6)
        r = p.add_run(name.upper() + " ")
        r.bold, r.font.size = True, Pt(8)
        self.runs(p, text, size=8)

    def save(self, path):
        self.doc.save(path)
        print("wrote", path, "| refs:", len(self.order), "| figs:", self.fig_n,
              "| tables:", self.tab_n, "| eqs:", self.eq_n)
